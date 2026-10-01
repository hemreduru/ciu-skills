#!/usr/bin/env python3
"""Usage: preview.py sunum.pptx <çıktı-klasörü> — soffice ile PDF ve slayt başına PNG üretir; soffice yoksa PREVIEW=none."""
import os
import shutil
import subprocess
import sys

from pptx import Presentation


def soffice(*args):
    return subprocess.run(["soffice", "--headless", *args], capture_output=True, text=True, timeout=300)


def main(pptx, out):
    if not shutil.which("soffice"):
        print("PREVIEW=none\nDETAIL=LibreOffice (soffice) kurulu değil; önizleme atlandı.")
        return 0
    base = os.path.splitext(os.path.basename(pptx))[0]
    os.makedirs(out, exist_ok=True)
    soffice("--convert-to", "pdf", "--outdir", out, pptx)
    pngs = []
    count = len(Presentation(pptx).slides)
    for n in range(1, count + 1):
        # soffice'in PNG dışa aktarımı yalnızca ilk slaytı verir; her slayt tek slaytlık kopyadan alınır.
        tmp = os.path.join(out, f".s{n}")
        os.makedirs(tmp, exist_ok=True)
        one = Presentation(pptx)
        lst = one.slides._sldIdLst
        for i in reversed(range(count)):
            if i != n - 1:
                one.part.drop_rel(lst[i].rId)
                lst.remove(lst[i])
        part = os.path.join(tmp, base + ".pptx")
        one.save(part)
        soffice("--convert-to", "png", "--outdir", tmp, part)
        src = os.path.join(tmp, base + ".png")
        if os.path.exists(src):
            dest = os.path.join(out, f"{base}-{n:02d}.png")
            os.replace(src, dest)
            pngs.append(dest)
        shutil.rmtree(tmp, ignore_errors=True)
    pdf = os.path.join(out, base + ".pdf")
    print(f"PREVIEW={'png' if pngs else 'pdf' if os.path.exists(pdf) else 'none'}\nDIR={out}")
    if os.path.exists(pdf):
        print(f"PDF={pdf}")
    print("\n".join(f"PNG={p}" for p in pngs))
    return 0 if pngs or os.path.exists(pdf) else 1


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Kullanım: preview.py sunum.pptx <çıktı-klasörü>")
        sys.exit(2)
    sys.exit(main(sys.argv[1], sys.argv[2]))
