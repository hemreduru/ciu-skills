import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build  # noqa: E402
import check  # noqa: E402
import common  # noqa: E402
import fetch_template  # noqa: E402

BG = '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:bgPr><a:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>'


def fake_template(path, example=("PRESENTATION TITLE", "TEXT ", "THANK YOU")):
    d = os.path.dirname(path)
    png = os.path.join(d, "bg.png")
    Image.new("RGB", (64, 36), (200, 30, 30)).save(png)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    layout = next(l for l in prs.slide_layouts if l.name == "Title Slide")
    for text in example:
        s = prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        _, rid = s.part.get_or_add_image_part(png)
        s._element.cSld.insert(0, etree.fromstring(BG % rid))
        s.shapes.add_textbox(Emu(1000000), Emu(1000000), Emu(5000000), Emu(500000)).text_frame.text = text
    prs.save(path)


def texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = self.tmp.name
        self.tpl = os.path.join(self.d, "sablon")
        os.makedirs(self.tpl)
        fake_template(os.path.join(self.tpl, "sablon-1.pptx"))
        self.img = os.path.join(self.d, "wide.png")
        Image.new("RGB", (400, 100), (10, 120, 200)).save(self.img)

    def tearDown(self):
        self.tmp.cleanup()

    def make(self, slides, template=1, **kw):
        deck = {"template": template, "lang": "tr", "slides": slides, **kw}
        p = os.path.join(self.d, "deck.json")
        with open(p, "w", encoding="utf-8") as f:
            json.dump(deck, f, ensure_ascii=False)
        out = os.path.join(self.d, "out")
        return build.build(p, self.tpl, self.d, out, os.path.join(self.d, "work"))

    def run_check(self, pptx):
        r = subprocess.run([sys.executable, os.path.join(HERE, "check.py"), pptx], capture_output=True, text=True)
        return r.returncode, r.stdout


GOOD = [
    {"layout": "title", "title": "Kampüs Tanıtımı", "subtitle": "Uluslararası Kıbrıs Üniversitesi", "notes": "Hoş geldiniz."},
    {"layout": "content", "title": "Neden UKÜ?", "bullets": ["Uluslararası kampüs", "Güçlü akademik kadro"], "notes": "Üç ana nokta."},
    {"layout": "closing", "title": "Teşekkürler", "subtitle": "www.ciu.edu.tr", "notes": "Sorular."},
]


class BuildTests(Base):
    def test_example_slides_removed_and_deck_clean(self):
        prs = Presentation(self.make(GOOD)["file"])
        self.assertEqual(len(prs.slides), 3)
        alltext = " ".join(t for s in prs.slides for t in texts(s))
        for bad in ("PRESENTATION TITLE", "THANK YOU", "TEXT"):
            self.assertNotIn(bad, alltext.replace("Teşekkürler", ""))
        self.assertIn("Kampüs Tanıtımı", alltext)
        self.assertIn("Teşekkürler", alltext)

    def test_placeholders_filled_unused_removed_notes_set(self):
        prs = Presentation(self.make(GOOD)["file"])
        for s in prs.slides:
            self.assertTrue(len(s.placeholders) >= 1)
            for ph in s.placeholders:
                self.assertTrue(ph.text_frame.text.strip(), ph.name)
        self.assertEqual(prs.slides[1].notes_slide.notes_text_frame.text, "Üç ana nokta.")

    def test_background_image_survives(self):
        prs = Presentation(self.make(GOOD)["file"])
        for s in prs.slides:
            blip = s._element.find(".//" + qn("a:blip"))
            self.assertIsNotNone(blip)
            self.assertIsNotNone(s.part.related_part(blip.get(qn("r:embed"))))

    def test_picture_keeps_aspect_ratio_inside_content_area(self):
        slides = GOOD[:1] + [{"layout": "picture", "title": "Kampüs", "bullets": ["Yeşil alan"], "image": self.img, "notes": "n"}]
        prs = Presentation(self.make(slides)["file"])
        pic = next(sh for sh in prs.slides[1].shapes if sh.shape_type == 13)
        self.assertAlmostEqual(pic.width / pic.height, 4.0, delta=0.05)
        self.assertLessEqual(pic.left + pic.width, prs.slide_width)

    def test_long_body_shrinks_font_but_not_below_minimum(self):
        bullets = ["Bu oldukça uzun bir madde metnidir ve kutuya sığması için küçülmelidir"] * 6
        prs = Presentation(self.make([{"layout": "content", "title": "Uzun", "bullets": bullets, "notes": "n"}])["file"])
        body = prs.slides[0].placeholders[1]
        size = body.text_frame.paragraphs[0].runs[0].font.size.pt
        self.assertLess(size, 24)
        self.assertGreaterEqual(size, 16)

    def test_user_template_without_profile_uses_its_own_layouts(self):
        own = os.path.join(self.d, "benim.pptx")
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "ÖRNEK"
        prs.save(own)
        r = self.make(GOOD, template=own)
        out = Presentation(r["file"])
        self.assertEqual(len(out.slides), 3)
        self.assertNotIn("ÖRNEK", " ".join(t for s in out.slides for t in texts(s)))

    def test_missing_title_is_rejected_in_turkish(self):
        with self.assertRaises(build.DeckError) as cm:
            self.make([{"layout": "content", "bullets": ["a"]}])
        self.assertIn("başlık", str(cm.exception))


class CheckTests(Base):
    def test_clean_deck_passes(self):
        code, out = self.run_check(self.make(GOOD)["file"])
        self.assertEqual(code, 0, out)
        self.assertEqual(out.splitlines()[0], "OK")

    def edit(self, fn, slides=GOOD):
        path = self.make(slides)["file"]
        prs = Presentation(path)
        fn(prs)
        prs.save(path)
        return path

    def test_overflow_is_error(self):
        def fn(prs):
            body = prs.slides[1].placeholders[1]
            body.text_frame.text = "Çok uzun metin " * 80
            body.text_frame.paragraphs[0].runs[0].font.size = Emu(25 * 12700)
        code, out = self.run_check(self.edit(fn))
        self.assertEqual(code, 1)
        self.assertIn("sığmıyor", out)

    def test_empty_placeholder_is_error(self):
        def fn(prs):
            prs.slides[1].placeholders[1].text_frame.text = ""
        code, out = self.run_check(self.edit(fn))
        self.assertEqual(code, 1)
        self.assertIn("boş", out)

    def test_leftover_example_text_is_error(self):
        def fn(prs):
            prs.slides[1].shapes.add_textbox(0, 0, 100, 100).text_frame.text = "PRESENTATION TITLE"
        code, out = self.run_check(self.edit(fn))
        self.assertEqual(code, 1)
        self.assertIn("örnek metin", out)

    def test_too_many_bullets_is_error(self):
        slides = [GOOD[0], {"layout": "content", "title": "Liste", "bullets": [f"Madde {i}" for i in range(9)], "notes": "n"}]
        code, out = self.run_check(self.make(slides)["file"])
        self.assertEqual(code, 1)
        self.assertIn("madde", out)

    def test_long_title_and_missing_notes_are_warnings_only(self):
        slides = [{"layout": "title", "title": "Bu başlık yedi kelimeden çok daha uzun olmuş durumda", "subtitle": "alt"}]
        code, out = self.run_check(self.make(slides)["file"])
        self.assertEqual(code, 0, out)
        self.assertIn("UYARI", out)
        self.assertIn("7 kelime", out)
        self.assertIn("konuşmacı notu", out)


    def test_example_text_in_table_group_and_notes_is_error(self):
        def table(prs):
            t = prs.slides[1].shapes.add_table(1, 1, 0, 0, 100000, 100000).table
            t.cell(0, 0).text = "Click to add text"
        def group(prs):
            g = prs.slides[1].shapes.add_group_shape()
            g.shapes.add_textbox(0, 0, 100, 100).text_frame.text = "Lorem ipsum"
        def notes(prs):
            prs.slides[1].notes_slide.notes_text_frame.text = "Lorem ipsum"
        for fn in (table, group, notes):
            code, out = self.run_check(self.edit(fn))
            self.assertEqual(code, 1, out)
            self.assertIn("örnek metin", out)

    def test_output_has_no_template_thumbnail_and_has_title(self):
        path = self.make(GOOD)["file"]
        with zipfile.ZipFile(path) as z:
            self.assertFalse([n for n in z.namelist() if "thumbnail" in n])
        self.assertEqual(Presentation(path).core_properties.title, "Kampüs Tanıtımı")


class UnitTests(unittest.TestCase):
    def test_estimate_grows_with_text_and_size(self):
        w = 8000000
        self.assertLess(common.est_height_pt(["kısa"], w, 20), common.est_height_pt(["uzun " * 100], w, 20))
        self.assertLess(common.est_height_pt(["metin " * 30], w, 16), common.est_height_pt(["metin " * 30], w, 28))

    def test_slug_turkish(self):
        self.assertEqual(common.slug("Öğrenci İşleri Tanıtımı!"), "ogrenci-isleri-tanitimi")

    def test_extract_templates_from_zip_with_turkish_folder(self):
        with tempfile.TemporaryDirectory() as d:
            z = os.path.join(d, "s.zip")
            with zipfile.ZipFile(z, "w") as f:
                for n in (1, 2, 3):
                    f.writestr(f"SUNUM-ŞABLONU/Sunum Template {n}.pptx", b"PK\x03\x04x")
            dest = os.path.join(d, "dest")
            fetch_template.extract_templates(z, dest)
            self.assertEqual(sorted(os.listdir(dest)), ["sablon-1.pptx", "sablon-2.pptx", "sablon-3.pptx"])

    def test_extract_templates_rejects_incomplete_zip(self):
        with tempfile.TemporaryDirectory() as d:
            z = os.path.join(d, "s.zip")
            with zipfile.ZipFile(z, "w") as f:
                f.writestr("Sunum Template 1.pptx", b"x")
            with self.assertRaises(RuntimeError):
                fetch_template.extract_templates(z, os.path.join(d, "dest"))


if __name__ == "__main__":
    unittest.main()
