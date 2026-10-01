#!/usr/bin/env python3
"""Usage: build.py deck.json --templates <klasör> --in <girdi> --out <çıktı> [--work <geçici>]
Deck JSON'dan şablonla .pptx üretir; şablonun örnek slaytları çıktıya girmez."""
import argparse
import copy
import json
import os
import shutil
import sys
import urllib.request

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import DATA, HERO_ROLES, ROLE_BG, pick_size, slug  # noqa: E402

TITLE_LADDER = {"hero": [44, 40, 36, 32, 28], "content": [36, 32, 28, 24]}
BODY = [24, 22, 20, 18, 16]
COLUMN = [24, 22, 20, 18, 16]
HEAD = [22, 20, 18]
SUB = [24, 22, 20, 18]


class DeckError(Exception):
    pass


def ph_by_idx(slide, idx):
    return next(p for p in slide.placeholders if p.placeholder_format.idx == idx)


def copy_bg(src, dst):
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is None:
        return
    new = copy.deepcopy(bg)
    for blip in new.iter(qn("a:blip")):
        part = src.part.related_part(blip.get(qn("r:embed")))
        blip.set(qn("r:embed"), dst.part.relate_to(part, RT.IMAGE))
    dst._element.cSld.insert(0, new)


def drop_slide(prs, i):
    lst = prs.slides._sldIdLst
    prs.part.drop_rel(lst[i].rId)
    lst.remove(lst[i])


def resolve_image(src, in_dir, work):
    os.makedirs(work, exist_ok=True)
    if src.startswith(("http://", "https://")):
        dest = os.path.join(work, slug(os.path.splitext(os.path.basename(src))[0]) + os.path.splitext(src)[1][:5])
        try:
            with urllib.request.urlopen(src, timeout=60) as r, open(dest, "wb") as f:
                shutil.copyfileobj(r, f)
        except Exception as e:
            raise DeckError(f"Görsel indirilemedi ({src}): {e}")
        src = dest
    elif not os.path.isabs(src):
        src = os.path.join(in_dir, src)
    if not os.path.isfile(src):
        raise DeckError(f"Görsel bulunamadı: {src}")
    try:
        im = Image.open(src)
        fixed = ImageOps.exif_transpose(im)
        if fixed is not im:
            src = os.path.join(work, slug(os.path.basename(src)) + ".png")
            fixed.save(src)
        im = fixed
    except Exception:
        raise DeckError(f"Görsel okunamadı: {src}")
    return src, im.size


def fill(shape, paras, size, color, font, lang, bold=False):
    tf = shape.text_frame
    tf.clear()
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        if bold:
            r.font.bold = True
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if font:
            r.font.name = font
        r._r.get_or_add_rPr().set("lang", lang)


def place(shape, rect, W, H):
    x0, y0, x1, y1 = rect
    shape.left, shape.top, shape.width, shape.height = Emu(int(x0 * W)), Emu(int(y0 * H)), Emu(int((x1 - x0) * W)), Emu(int((y1 - y0) * H))


def lines(v):
    return [v] if isinstance(v, str) else [str(x) for x in (v or [])]


def add_picture(slide, path, size, box, alt):
    bx, by, bw, bh = box
    k = min(bw / size[0], bh / size[1])
    w, h = int(size[0] * k), int(size[1] * k)
    pic = slide.shapes.add_picture(path, Emu(int(bx + (bw - w) / 2)), Emu(int(by)), Emu(w), Emu(h))
    pic._element.nvPicPr.cNvPr.set("descr", alt)


def build(deck_path, templates_dir, in_dir, out_dir, work):
    with open(deck_path, encoding="utf-8") as f:
        deck = json.load(f)
    slides = deck.get("slides") or []
    if not slides:
        raise DeckError("Sunumda slayt yok.")
    lang = {"en": "en-US"}.get(deck.get("lang", "tr"), "tr-TR")
    tpl = str(deck.get("template", "1"))
    profile = DATA["templates"].get(tpl)
    if profile:
        path = os.path.join(templates_dir, f"sablon-{tpl}.pptx")
        layouts, font, colors = DATA["layouts"], DATA["font"], DATA["colors"]
    else:
        path = tpl if os.path.isabs(tpl) else os.path.join(in_dir, tpl)
        layouts, font, colors = DATA["generic_layouts"], None, {}
    if not os.path.isfile(path):
        raise DeckError(f"Şablon dosyası bulunamadı: {path}")
    prs = Presentation(path)
    W, H = prs.slide_width, prs.slide_height
    original = len(prs.slides._sldIdLst)
    by_name = {l.name: l for l in prs.slide_layouts}
    sources = list(prs.slides)

    for n, spec in enumerate(slides, 1):
        role = spec.get("layout", "content")
        if role not in layouts:
            raise DeckError(f"Slayt {n}: bilinmeyen düzen '{role}'. Kullanılabilir: {', '.join(layouts)}.")
        if not str(spec.get("title", "")).strip():
            raise DeckError(f"Slayt {n}: başlık boş olamaz.")
        layout = by_name.get(layouts[role])
        if layout is None:
            raise DeckError(f"Şablonda '{layouts[role]}' düzeni yok. Mevcut düzenler: {', '.join(by_name)}.")
        slide = prs.slides.add_slide(layout)
        ids = DATA["placeholders"][layouts[role]]
        ph = {k: ph_by_idx(slide, i) for k, i in ids.items() if any(p.placeholder_format.idx == i for p in slide.placeholders)}
        hero = role in HERO_ROLES
        bgkey = ROLE_BG.get(role, "content")
        dark = bool(profile) and profile["tone"][bgkey] == "dark"
        c_title = colors.get("white" if dark else "wine")
        c_body = colors.get("white" if dark else "ink")
        if profile:
            copy_bg(sources[profile["bg"][bgkey]], slide)
        title = str(spec["title"]).strip()
        sub = lines(spec.get("subtitle"))
        bullets = lines(spec.get("bullets"))
        alt = spec.get("alt") or title

        def put(key, paras, ladder, color, bullets_=False, bold=False, rect=None):
            shape = ph[key]
            if rect and profile:
                place(shape, rect, W, H)
            if not paras:
                shape._element.getparent().remove(shape._element)
                return
            size = pick_size(paras, shape.width, shape.height, ladder, bullets_, bold)
            fill(shape, paras, size, color, font, lang, bold)

        if profile and hero:
            x0, y0, x1, y1 = profile["hero"]
            mid = y0 + (y1 - y0) * 0.55
            put("title", [title], TITLE_LADDER["hero"], c_title, bold=True, rect=[x0, y0, x1, mid])
            put("subtitle", sub, SUB, c_body, rect=[x0, mid + 0.02, x1, y1])
        elif hero:
            put("title", [title], TITLE_LADDER["hero"], c_title, bold=bool(profile))
            put("subtitle", sub, SUB, c_body)
        else:
            sx0, sy0, sx1, sy1 = profile["safe"] if profile else (0, 0, 1, 1)
            ty1 = sy0 + 0.13
            by0, gap = ty1 + 0.02, 0.03
            colw = (sx1 - sx0 - gap) / 2
            put("title", [title], TITLE_LADDER["content"], c_title, bold=True, rect=[sx0, sy0, sx1, ty1])
            if role == "content":
                put("body", bullets, BODY, c_body, True, rect=[sx0, by0, sx1, sy1])
            elif role == "picture":
                img = spec.get("image")
                if not img:
                    raise DeckError(f"Slayt {n}: 'picture' düzeni için image gerekir.")
                path_, size = resolve_image(img, in_dir, work)
                if "picture" in ph:
                    pp = ph["picture"]
                    box = (pp.left, pp.top, pp.width, pp.height)
                    pp._element.getparent().remove(pp._element)
                    put("body", bullets, COLUMN, c_body)
                else:
                    split = sx0 + (sx1 - sx0) * 0.44
                    box = (int((split + gap) * W), int(by0 * H), int((sx1 - split - gap) * W), int((sy1 - by0) * H))
                    put("body", bullets, COLUMN, c_body, True, rect=[sx0, by0, split, sy1])
                add_picture(slide, path_, size, box, alt)
            elif role == "title_only":
                if spec.get("image"):
                    path_, size = resolve_image(spec["image"], in_dir, work)
                    add_picture(slide, path_, size, (int(sx0 * W), int(by0 * H), int((sx1 - sx0) * W), int((sy1 - by0) * H)), alt)
            elif role == "two":
                put("left", lines(spec.get("left")), COLUMN, c_body, True, rect=[sx0, by0, sx0 + colw, sy1])
                put("right", lines(spec.get("right")), COLUMN, c_body, True, rect=[sx1 - colw, by0, sx1, sy1])
            elif role == "comparison":
                hy1 = by0 + 0.08
                put("left_title", lines(spec.get("left_title")), HEAD, c_title, bold=True, rect=[sx0, by0, sx0 + colw, hy1])
                put("right_title", lines(spec.get("right_title")), HEAD, c_title, bold=True, rect=[sx1 - colw, by0, sx1, hy1])
                put("left", lines(spec.get("left")), COLUMN, c_body, True, rect=[sx0, hy1 + 0.01, sx0 + colw, sy1])
                put("right", lines(spec.get("right")), COLUMN, c_body, True, rect=[sx1 - colw, hy1 + 0.01, sx1, sy1])
        for p in list(slide.placeholders):
            if p.has_text_frame and not p.text_frame.text.strip():
                p._element.getparent().remove(p._element)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])

    for _ in range(original):
        drop_slide(prs, 0)
    pkg = prs.part.package
    for rid, rel in list(pkg._rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            pkg._rels.pop(rid)
    prs.core_properties.title = str(slides[0]["title"]).strip()
    os.makedirs(out_dir, exist_ok=True)
    name = slug(deck.get("name") or slides[0]["title"])
    out = os.path.join(out_dir, f"{name}.pptx")
    prs.save(out)
    return {"file": out, "slides": len(slides)}


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--templates", required=True)
    ap.add_argument("--in", dest="in_dir", default=".")
    ap.add_argument("--out", required=True)
    ap.add_argument("--work", default="/tmp/ciu-slides-work")
    a = ap.parse_args()
    try:
        r = build(a.deck, a.templates, a.in_dir, a.out, a.work)
    except DeckError as e:
        print(f"HATA: {e}")
        sys.exit(1)
    print(f"FILE={r['file']}\nSLIDES={r['slides']}\nSıradaki adım: check.py ile denetle.")
